import re

import arrow
import openpyxl

from app_setup import db
from models.mixins import FilterMixin
from modules.tickets.models import TicketModel


class TicketsService(FilterMixin):

    @classmethod
    def get_tickets(cls, filters=None):
        query = TicketModel.query
        # query = query.order_by(TicketModel.id.desc())
        query = query.order_by(
            TicketModel.hotline_ticket_date.asc(),
            TicketModel.hotline_ticket_no.asc()
        )
        query = cls.apply_filters(query, TicketModel, filters)

        return query.all()

    @classmethod
    def get_tickets_dicts(cls, filters=None):
        tickets = cls.get_tickets(filters)
        return [t.to_python() for t in tickets]

    @classmethod
    def get_ticket_dict(cls, ticket_id):
        ticket = TicketModel.query.get(ticket_id)
        if not ticket:
            return None
        ticket_dict = ticket.to_python()

        return ticket_dict

    @classmethod
    def create(cls, payload):
        ticket = TicketModel()

        ticket.unit_id = payload.get('unit_id')
        ticket.status_id = payload.get('status_id')
        ticket.docs_status_id = payload.get('docs_status_id')
        ticket.subj = payload.get('subj')
        ticket.body = payload.get('body')
        ticket.reason = payload.get('reason')
        ticket.needs_to_be_done = payload.get('needs_to_be_done')
        ticket.comleted_work = payload.get('comleted_work')
        ticket.is_hotline = payload.get('is_hotline')
        ticket.hotline_ticket_no = payload.get('hotline_ticket_no')
        ticket.hotline_ticket_date = payload.get('hotline_ticket_date')
        ticket.hotline_ticket_json = payload.get('hotline_ticket_json')
        ticket.hotline_fetch_date = payload.get('hotline_fetch_date')
        ticket.iac_ticket_no = payload.get('iac_ticket_no')
        ticket.ticket_comments_json = payload.get('ticket_comments_json')

        db.session.add(ticket)
        db.session.commit()

        return ticket

    @classmethod
    def get_tickets_from_excel(cls, xlsx_path,
                               iterion_tickets_xlsx_filename,
                               iterion_tickets_xlsx_tab_no,
                               iterion_header_rows,
                               iac_tickets_xlsx_filename,
                               iac_tickets_header_rows,
                               iac_tickets_year):

        workbook = openpyxl.load_workbook(
            f'{xlsx_path}{iterion_tickets_xlsx_filename}', data_only=True)

        worksheet = workbook.worksheets[iterion_tickets_xlsx_tab_no]
        max_row = worksheet.max_row
        max_col = worksheet.max_column

        tickets_map = cls.get_tickets_hotline_iac_map(
            iac_tickets_xlsx_path=xlsx_path,
            iac_tickets_xlsx_filename=iac_tickets_xlsx_filename,
            header_rows_to_skip=iac_tickets_header_rows,
            iac_tickets_year=iac_tickets_year
        )

        tickets_from_excel = []
        print('Source file: {} (worksheet #{} {})'.format(
            iterion_tickets_xlsx_filename,
            iterion_tickets_xlsx_tab_no,
            workbook.sheetnames[iterion_tickets_xlsx_tab_no]
        ))
        for row in range(iterion_header_rows + 1, max_row + 1):
            cell_a = worksheet.cell(row, 1).value
            hotline_ticket_no = worksheet.cell(row, 2).value
            hotline_ticket_date = worksheet.cell(row, 3).value
            subj = worksheet.cell(row, 4).value
            body = worksheet.cell(row, 5).value
            try:
                hotline_ticket_dt = arrow.get(
                    hotline_ticket_date,
                    'DD-MM-YYYY'
                ).datetime
            except Exception as e:
                print('ERROR Date parse: {} - {} - {}'.format(
                    worksheet.cell(row, 3).value, subj, e))

            cell_f = worksheet.cell(row, 6).value
            comleted_work = worksheet.cell(row, 7).value
            cell_h = worksheet.cell(row, 8).value
            cell_i = worksheet.cell(row, 9).value
            cell_j = worksheet.cell(row, 10).value
            cell_k = worksheet.cell(row, 11).value
            cell_l = worksheet.cell(row, 12).value
            cell_m = worksheet.cell(row, 13).value
            cell_n = worksheet.cell(row, 14).value
            cell_o = worksheet.cell(row, 15).value

            comments = dict(
                c_a=str(cell_a),
                c_f=str(cell_f),
                c_h=str(cell_h),
                c_i=str(cell_i),
                c_j=str(cell_j),
                c_k=str(cell_k),
                c_l=str(cell_l),
                c_m=str(cell_m),
                c_n=str(cell_n),
                c_o=str(cell_o),
            )

            ticket = dict(
                hotline_ticket_no=hotline_ticket_no,
                hotline_ticket_date=str(hotline_ticket_date),
                iac_ticket_no=tickets_map.get(str(hotline_ticket_no)),
                subj=subj,
                body=body,
                comleted_work=comleted_work,
                comments=comments,
            )
            tickets_from_excel.append(ticket)

        return tickets_from_excel

    @staticmethod
    def parse_tickets_no_pair(tickets_no_str: str, iac_tickets_year: str) -> (str, str):
        """
        Parse ticket numbers from string with re

        :param tickets_no_str: e.g. `в филиале- 1124/2023в службе поддержки ИАЦ-1205733`
        :param iac_tickets_year: year of iac tickets? e.g. 2023
        :return: a pair of ticket numbers
        """

        # search ticket_no here '1124/2023в'
        match = re.search(r'\d+/{}'.format(iac_tickets_year), tickets_no_str)
        iac_ticket_no = match[0][:-5] if match else None

        # search ticket_no here 'ИАЦ-1205733'
        match = re.search(r'ИАЦ-\d+', tickets_no_str)
        hotline_ticket_no = match[0][4:] if match else None

        return hotline_ticket_no, iac_ticket_no

    @classmethod
    def get_tickets_hotline_iac_map(
            cls, iac_tickets_xlsx_path: str, iac_tickets_xlsx_filename: str,
            header_rows_to_skip: int, iac_tickets_year: str) -> dict:
        """
        Parse xlsx w tickets list from legacy IS IAC Chelyabinsk

        :param iac_tickets_xlsx_path: path to xlsx xlsx_filename
        :param iac_tickets_xlsx_filename: xlsx w tickets list
        :param header_rows_to_skip: qty of headers row in excel file
        :param iac_tickets_year: year of iac tickets
        :return: dict {hotline_ticket_no: iac_ticket_no}
        """
        tickets_map = {}
        workbook = openpyxl.load_workbook(
            f'{iac_tickets_xlsx_path}{iac_tickets_xlsx_filename}', data_only=True)
        worksheet = workbook.worksheets[0]
        max_row = worksheet.max_row

        for row in range(header_rows_to_skip + 1, max_row + 1):
            # row_pp = worksheet.cell(row, 1).value
            tickets_no_str = worksheet.cell(row, 3).value
            hotline_ticket_no, iac_ticket_no = cls.parse_tickets_no_pair(tickets_no_str, iac_tickets_year)
            tickets_map[hotline_ticket_no] = iac_ticket_no
            # print(f'{row_pp} - {tickets_no_str} - {iac_ticket_no}/{hotline_ticket_no}')

        return tickets_map
