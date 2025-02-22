#
# Load Tomsk hardware models list from xlsx
#

import openpyxl

from app_setup import db, create_app
from modules.hwdb.models import HwTomskModels

XLSX_PATH = '../dumps/'
XLSX_FILENAME = 'tomsk_models_manuf_29012025.xlsx'
HEADER_ROW = 4


def main():
    workbook = openpyxl.load_workbook(f'{XLSX_PATH}{XLSX_FILENAME}', data_only=True)
    worksheet = workbook.worksheets[0]
    max_row = worksheet.max_row
    max_col = worksheet.max_column

    print('Source file: {}'.format(XLSX_FILENAME))

    app = create_app()
    with app.app_context():
        for row in range(HEADER_ROW + 1, max_row + 1):
            line_num = worksheet.cell(row, 1).value
            hw_type = worksheet.cell(row, 2).value
            manuf = worksheet.cell(row, 3).value
            model = worksheet.cell(row, 4).value

            print(f'{line_num} - {hw_type} - {manuf} - {model}')
            unit = HwTomskModels(
                line_num=line_num,
                hw_type=hw_type,
                manuf=manuf,
                model=model
            )
            db.session.add(unit)
            db.session.commit()


if __name__ == '__main__':
    main()
