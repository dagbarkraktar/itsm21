#
# One time script: Load hwdb units from xlsx file
#

import openpyxl

from app_setup import db, create_app
from modules.hwdb.models import HwDbModel

XLSX_PATH = '../dumps/'
# XLSX_FILENAME = 'HWDB_IMPORT_2023_ARM.xlsx'
# XLSX_FILENAME = 'HWDB_IMPORT_2025_ARM_PRINT_MFU.xlsx'
XLSX_FILENAME = ''
HEADER_ROW = 1


def main():
    workbook = openpyxl.load_workbook(f'{XLSX_PATH}{XLSX_FILENAME}', data_only=True)
    worksheet = workbook.worksheets[0]
    max_row = worksheet.max_row
    max_col = worksheet.max_column

    print('Source file: {}'.format(XLSX_FILENAME))

    app = create_app()
    with app.app_context():
        for row in range(HEADER_ROW + 1, max_row + 1):
            invnum = worksheet.cell(row, 1).value
            if not invnum:
                continue
            type_id = worksheet.cell(row, 2).value
            legacy_user = worksheet.cell(row, 3).value
            empl_id = worksheet.cell(row, 4).value
            location = worksheet.cell(row, 5).value
            status_id = worksheet.cell(row, 6).value
            manuf = worksheet.cell(row, 7).value
            model = worksheet.cell(row, 8).value
            vendor_id = worksheet.cell(row, 9).value
            serialnum = worksheet.cell(row, 10).value
            year = worksheet.cell(row, 11).value
            warranty = worksheet.cell(row, 12).value
            accounting = worksheet.cell(row, 13).value
            comments = worksheet.cell(row, 14).value
            buhtext = worksheet.cell(row, 15).value
            buh_os = worksheet.cell(row, 16).value
            check = worksheet.cell(row, 17).value
            gas_id = worksheet.cell(row, 18).value
            last_maintenance_id = worksheet.cell(row, 19).value

            unit = HwDbModel()
            unit.invnum = invnum
            unit.type_id = type_id
            unit.legacy_user = legacy_user
            unit.empl_id = empl_id
            unit.location = location
            unit.status_id = status_id
            unit.manuf = manuf
            unit.model = model
            unit.vendor_id = vendor_id
            unit.serialnum = serialnum
            unit.year = year
            unit.warranty = warranty
            unit.accounting = accounting
            unit.comments = comments
            unit.buhtext = buhtext
            unit.buh_os = buh_os
            unit.check = check
            unit.gas_id = gas_id
            unit.last_maintenance_id = last_maintenance_id

            db.session.add(unit)
            db.session.commit()

        print('Processed: {} lines'.format(row))


if __name__ == '__main__':
    main()
