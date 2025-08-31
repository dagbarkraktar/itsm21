#
# Periodic usage: Load iac_ticket_no list from xlsx and match it with hotline tickets in db
#
import argparse
import os
import sys
import openpyxl

from app_setup import db, create_app
from modules.tickets.models import TicketModel
from modules.tickets.services import TicketsService


def parse_arguments():
    """Parse command line arguments."""
    parser = argparse.ArgumentParser(
        description='Load IAC ticket numbers from xlsx and match with hotline tickets in database',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='''
Examples:
  %(prog)s tickets_is_iac_2025-03-21.xlsx
  %(prog)s tickets_is_iac_2025-03-21.xlsx --header 5 --year 2025
  %(prog)s /path/to/tickets.xlsx --no-save-to-db
  %(prog)s tickets.xlsx -H 2 -y 2024 --verbose
        '''
    )

    parser.add_argument(
        'input_file',
        help='Path to the input xlsx file with IAC tickets'
    )

    parser.add_argument(
        '-H', '--header',
        type=int,
        default=3,
        help='Number of header rows to skip (default: 3)'
    )

    parser.add_argument(
        '-y', '--year',
        default='2025',
        help='IAC tickets year (default: 2025)'
    )

    parser.add_argument(
        '--save-to-db',
        action='store_true',
        default=True,
        help='Save results to database (default: True)'
    )

    parser.add_argument(
        '--no-save-to-db',
        dest='save_to_db',
        action='store_false',
        help='Do not save results to database'
    )

    parser.add_argument(
        '-v', '--verbose',
        action='store_true',
        help='Enable verbose output'
    )

    return parser.parse_args()


def main():
    args = parse_arguments()

    # Validate input file
    if not os.path.isfile(args.input_file):
        print(f"Error: Input file '{args.input_file}' does not exist.", file=sys.stderr)
        sys.exit(1)

    # Extract path and filename from input file
    xlsx_path = os.path.dirname(args.input_file) + '/' if os.path.dirname(args.input_file) else ''
    xlsx_filename = os.path.basename(args.input_file)

    if args.verbose:
        print(f"Input file: {args.input_file}")
        print(f"Header rows to skip: {args.header}")
        print(f"IAC tickets year: {args.year}")
        print(f"Save to database: {args.save_to_db}")

    app = create_app()

    try:
        workbook = openpyxl.load_workbook(args.input_file, data_only=True)
    except Exception as e:
        print(f"Error: Cannot open input file '{args.input_file}': {e}", file=sys.stderr)
        sys.exit(1)

    worksheet = workbook.worksheets[0]
    max_row = worksheet.max_row
    max_col = worksheet.max_column

    # preview in console
    print(f'Processing file: {xlsx_filename}')
    processed_count = 0
    for row in range(args.header + 1, max_row + 1):
        row_pp = worksheet.cell(row, 1).value
        tickets_no_str = worksheet.cell(row, 3).value

        if not tickets_no_str:
            continue

        hotline_ticket_no, iac_ticket_no = TicketsService.parse_tickets_no_pair(
            tickets_no_str, args.year)

        if args.verbose:
            print(f'{row_pp} - {tickets_no_str} - {iac_ticket_no}/{hotline_ticket_no}')

        processed_count += 1

    print(f'Processed {processed_count} ticket entries')

    # write iac_ticket_no to db
    if args.save_to_db:
        print('Starting database update...')
        try:
            with app.app_context():
                tickets_map = TicketsService.get_tickets_hotline_iac_map(
                    iac_tickets_xlsx_path=xlsx_path,
                    iac_tickets_xlsx_filename=xlsx_filename,
                    header_rows_to_skip=args.header,
                    iac_tickets_year=args.year
                )

                updated_count = 0
                skipped_count = 0

                for hotline_ticket_no, iac_ticket_no in tickets_map.items():
                    ticket = TicketModel.query.filter(
                        TicketModel.hotline_ticket_no == hotline_ticket_no
                    ).first()
                    if not ticket:
                        if args.verbose:
                            print(f'SKIP LINE ({hotline_ticket_no}) - ticket not found in database')
                        skipped_count += 1
                        continue
                    ticket.iac_ticket_no = iac_ticket_no
                    db.session.commit()
                    updated_count += 1

                print(f'Database update completed: {updated_count} tickets updated, {skipped_count} skipped')
        except Exception as e:
            print(f"Error during database update: {e}", file=sys.stderr)
            sys.exit(1)
    else:
        print('Database update skipped (--no-save-to-db flag used)')


if __name__ == '__main__':
    try:
        main()
    except KeyboardInterrupt:
        print("\nOperation cancelled by user.", file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(f"Unexpected error: {e}", file=sys.stderr)
        sys.exit(1)
