#!/usr/bin/env python3
#
# On demand script: Convert 1C report from IAC to simple readable xlsx table
#

import argparse
import os
import sys
import openpyxl
import xlsxwriter


def parse_arguments():
    """Parse command line arguments."""
    parser = argparse.ArgumentParser(
        description='Convert 1C report to readable xlsx table',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='''
Examples:
  %(prog)s input.xlsx
  %(prog)s input.xlsx --header-rows 50
  %(prog)s /path/to/file.xlsx -r 30
        '''
    )

    parser.add_argument(
        'input_file',
        help='Path to the input xlsx file'
    )

    parser.add_argument(
        '-r', '--header-rows',
        type=int,
        default=40,
        help='Number of header rows to skip (default: 40)'
    )

    parser.add_argument(
        '-o', '--output',
        help='Output file path (default: input_file_readable.xlsx)'
    )

    parser.add_argument(
        '-v', '--verbose',
        action='store_true',
        help='Enable verbose output'
    )

    return parser.parse_args()


def main():
    args = parse_arguments()

    # Validate input file
    if not os.path.isfile(args.input_file):
        print(f"Error: Input file '{args.input_file}' does not exist.", file=sys.stderr)
        sys.exit(1)

    # Generate output filename if not provided
    if args.output:
        output_file = args.output
    else:
        base_name = os.path.splitext(args.input_file)[0]
        output_file = f"{base_name}_readable.xlsx"

    if args.verbose:
        print(f"Input file: {args.input_file}")
        print(f"Output file: {output_file}")
        print(f"Header rows to skip: {args.header_rows}")

    try:
        workbook = openpyxl.load_workbook(args.input_file, data_only=True)
    except Exception as e:
        print(f"Error: Cannot open input file '{args.input_file}': {e}", file=sys.stderr)
        sys.exit(1)

    worksheet = workbook.worksheets[0]
    max_row = worksheet.max_row

    try:
        output_workbook = xlsxwriter.Workbook(output_file)
        output_worksheet = output_workbook.add_worksheet()
    except Exception as e:
        print(f"Error: Cannot create output file '{output_file}': {e}", file=sys.stderr)
        sys.exit(1)

    print(f'Processing file: {os.path.basename(args.input_file)}')

    cnt = 0
    row = 0
    out_cur_row = 0
    for row in range(args.header_rows + 1, max_row + 1):
        line_no = worksheet.cell(row, 1).value
        if not line_no or not isinstance(line_no, int):
            continue
        cnt += 1

        name = worksheet.cell(row, 4).value
        invnum = worksheet.cell(row, 11).value
        unit = worksheet.cell(row, 19).value
        price = worksheet.cell(row, 22).value
        qty = worksheet.cell(row, 27).value
        price_sum = worksheet.cell(row, 34).value

        if args.verbose:
            print('#{} - {} - {} - {} - {} - {}'.format(
                line_no,
                invnum,
                unit,
                price,
                qty,
                price_sum
            ))

        # write line to output xlsx
        # TODO: add cell formats (text, date, currency etc)
        output_worksheet.write(out_cur_row, 0, line_no)
        output_worksheet.write(out_cur_row, 1, name)
        output_worksheet.write(out_cur_row, 2, invnum or '')
        output_worksheet.write(out_cur_row, 3, unit)
        output_worksheet.write(out_cur_row, 4, price)
        output_worksheet.write(out_cur_row, 5, qty)
        output_worksheet.write(out_cur_row, 6, price_sum)
        out_cur_row += 1

    print('Processed: {} data lines (raw: {})'.format(cnt, row))
    print(f'Output saved to: {output_file}')

    try:
        output_workbook.close()
    except Exception as e:
        print(f"Error: Cannot save output file '{output_file}': {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == '__main__':
    try:
        main()
    except KeyboardInterrupt:
        print("\nOperation cancelled by user.", file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(f"Unexpected error: {e}", file=sys.stderr)
        sys.exit(1)
